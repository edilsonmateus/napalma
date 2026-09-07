from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("Manual_de_Cadastro_e_Operacao_Inicial_77Gira_para_Casas.docx")

NAVY = "10233F"
ORANGE = "F28C28"
BLUE = "2463EB"
GRAY = "64748B"
LIGHT = "E2E8F0"
PALE = "F8FAFC"
GREEN = "0F766E"
RED = "B42318"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, edge_data in edges.items():
        tag = "w:" + edge_name
        edge = borders.find(qn(tag))
        if edge is None:
            edge = OxmlElement(tag)
            borders.append(edge)
        for key, value in edge_data.items():
            edge.set(qn("w:" + key), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_cell_width(cell, width):
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_width = tc_pr.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        tc_pr.append(tc_width)
    tc_width.set(qn("w:w"), str(int(width.twips)))
    tc_width.set(qn("w:type"), "dxa")


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_keep_together(paragraph, value=True):
    paragraph.paragraph_format.keep_together = value


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_run(paragraph, text, bold=False, italic=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_body(doc, text, before=0, after=5, keep=False):
    p = doc.add_paragraph(style="Corpo")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    set_keep_together(p, keep)
    p.add_run(text)
    return p


def add_rich_body(doc, parts, before=0, after=5, keep=False):
    p = doc.add_paragraph(style="Corpo")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    set_keep_together(p, keep)
    for part in parts:
        if isinstance(part, str):
            p.add_run(part)
        else:
            add_run(p, part.get("text", ""), part.get("bold", False), part.get("italic", False), part.get("color"))
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="Lista")
    p.paragraph_format.left_indent = Inches(0.22 + level * 0.20)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("• ")
    p.add_run(text)
    return p


def add_numbered_step(doc, number, title, details):
    p = doc.add_paragraph(style="Passo")
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    set_keep_with_next(p)
    add_run(p, f"{number}. ", bold=True, color=ORANGE)
    add_run(p, title, bold=True, color=NAVY)
    if details:
        add_body(doc, details, after=4)
    return p


def add_note(doc, label, text, color=BLUE):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(0.08)
    table.columns[1].width = Inches(6.75)
    set_cell_shading(table.cell(0, 0), color)
    cell = table.cell(0, 1)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, PALE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    add_run(p, f"{label}  ", bold=True, color=color)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_image_placeholder(doc, number, title, capture, crop="Mostre a tela inteira, sem dados sensíveis reais.", height=1.55):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.rows[0].height = Inches(height)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = table.cell(0, 0)
    cell.width = Inches(6.82)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "FFFFFF")
    border = {"val": "dashed", "sz": "10", "color": "94A3B8", "space": "4"}
    set_cell_border(cell, top=border, bottom=border, left=border, right=border)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    add_run(p, f"INSERIR IMAGEM {number:02d}", bold=True, color=BLUE, size=11)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(3)
    add_run(p2, title, bold=True, color=NAVY)
    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(1)
    add_run(p3, "Capturar: ", bold=True, color=GRAY, size=8.5)
    add_run(p3, capture, color=GRAY, size=8.5)
    p4 = cell.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_after = Pt(10)
    add_run(p4, "Recorte: ", bold=True, color=GRAY, size=8.5)
    add_run(p4, crop, color=GRAY, size=8.5)
    cap = doc.add_paragraph(style="Legenda")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Imagem {number:02d} — {title}")
    return table


def add_field_table(doc, rows, widths=(1.55, 2.15, 3.10)):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(("Campo", "Como preencher", "Cuidados")):
        set_cell_shading(hdr.cells[i], NAVY)
        p = hdr.cells[i].paragraphs[0]
        add_run(p, text, bold=True, color="FFFFFF", size=9)
    for field, how, care in rows:
        cells = table.add_row().cells
        for i, text in enumerate((field, how, care)):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_run(p, text, bold=(i == 0), color=NAVY if i == 0 else None, size=8.5)
            border = {"val": "single", "sz": "4", "color": LIGHT}
            set_cell_border(cells[i], top=border, bottom=border, left=border, right=border)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_checklist(doc, items):
    for item in items:
        p = doc.add_paragraph(style="Lista")
        p.paragraph_format.left_indent = Inches(0.08)
        p.paragraph_format.space_after = Pt(3)
        add_run(p, "☐  ", color=BLUE, size=12)
        p.add_run(item)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, bold, color, before, after in [
        ("Title", 27, True, NAVY, 0, 12),
        ("Heading 1", 18, True, "111827", 14, 8),
        ("Heading 2", 13, True, "111827", 11, 5),
        ("Heading 3", 11, True, "111827", 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    # The default Word theme may attach a bottom rule to the Title style.
    # Remove it explicitly so the cover title follows the 77Gira visual system.
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name in ["Corpo", "Lista", "Passo", "Legenda"]:
        if style_name not in styles:
            styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    body = styles["Corpo"]
    body.font.name = "Arial"
    body.font.size = Pt(10)
    body.font.color.rgb = RGBColor.from_string(NAVY)
    body.paragraph_format.line_spacing = 1.14
    body.paragraph_format.space_after = Pt(5)
    lst = styles["Lista"]
    lst.font.name = "Arial"
    lst.font.size = Pt(9.7)
    lst.font.color.rgb = RGBColor.from_string(NAVY)
    step = styles["Passo"]
    step.font.name = "Arial"
    step.font.size = Pt(10.5)
    legend = styles["Legenda"]
    legend.font.name = "Arial"
    legend.font.size = Pt(8)
    legend.font.italic = True
    legend.font.color.rgb = RGBColor.from_string(GRAY)

    for sec in doc.sections:
        footer = sec.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_run(p, "77Gira  |  Manual para casas  |  versão 1.1 — 07/09/2026", color=GRAY, size=8)
        add_page_number(footer.add_paragraph())

    props = doc.core_properties
    props.title = "Manual de cadastro e operação inicial do 77Gira para casas"
    props.subject = "Conta, casa, produtores, eventos e Cardápio Essencial"
    props.author = "77Gira"
    props.keywords = "77Gira, casas, produtores, eventos, cardápio, manual"

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Inches(0.72)
    add_run(p, "77GIRA", bold=True, color=ORANGE, size=14)
    title = doc.add_paragraph(style="Title")
    title.add_run("Manual de cadastro e operação inicial")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    add_run(sub, "Guia passo a passo para casas, gestores e produtores", color=GRAY, size=14)
    add_body(doc, "Este manual acompanha a casa desde a criação da conta até a publicação do primeiro evento e do Cardápio Essencial. Os nomes dos botões e campos seguem a interface atual do 77Gira.", after=14)
    add_note(doc, "PARA QUEM É", "Proprietários, representantes, gestores autorizados e produtores responsáveis pela agenda da casa.", GREEN)
    add_note(doc, "IMPORTANTE", "Cada pessoa deve usar sua própria conta. Nunca compartilhe senha, código recebido por e-mail ou acesso entre integrantes da equipe.", RED)
    add_body(doc, "Versão 1.1  |  Atualizada em 7 de setembro de 2026", before=18, after=0)
    doc.add_page_break()

    add_heading(doc, "Como usar este manual", 1)
    add_body(doc, "Siga os capítulos na ordem indicada no primeiro acesso. Depois da implantação, use as seções de evento e cardápio como consulta operacional.")
    add_heading(doc, "Mapa do processo", 2)
    for item in [
        "Criar uma conta pessoal no 77Gira.",
        "Localizar a casa e solicitar acesso à gestão. Se ela ainda não existir, enviar uma solicitação interna de inclusão e guardar o protocolo.",
        "Aguardar a análise e, quando solicitado, ler e assinar o documento formal.",
        "Abrir o Hub de Gestão e conferir os dados da casa.",
        "Cadastrar ou vincular os produtores responsáveis.",
        "Criar, revisar e publicar os eventos.",
        "Cadastrar, organizar e publicar o Cardápio Essencial.",
        "Fazer a checagem final no perfil público e no Explorar.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Antes de começar", 2)
    add_checklist(doc, [
        "Separe um e-mail individual que você consulta e uma senha exclusiva para o 77Gira.",
        "Tenha em mãos nome, endereço, bairro, cidade, UF, dias de funcionamento, telefone, Instagram e uma boa imagem da casa.",
        "Confirme quem é o responsável legal ou autorizado pela solicitação de acesso.",
        "Separe CPF ou CNPJ do solicitante e uma explicação objetiva de como o vínculo com a casa pode ser confirmado.",
        "Para o evento, tenha título, descrição, data e hora de início e fim, preço, link de ingresso e cartaz 4:5 quando houver.",
        "Para o cardápio, prepare uma seleção curta de itens, categorias, descrições e preços.",
    ])
    add_note(doc, "DADOS PESSOAIS", "Ao produzir as capturas para este manual, oculte e-mails, documentos, telefones, senhas e códigos reais.", RED)

    # Account creation
    doc.add_page_break()
    add_heading(doc, "1  Criação da conta", 1)
    add_body(doc, "A conta é pessoal. Criá-la não transforma automaticamente o usuário em gestor de uma casa; o acesso profissional é solicitado em uma etapa separada.")
    add_numbered_step(doc, 1, "Abra o 77Gira", "Na tela inicial ou em Configurações, escolha Criar conta.")
    add_image_placeholder(doc, 1, "Entrada para criação de conta", "Tela inicial ou Configurações com o botão “Criar conta” destacado.")
    add_numbered_step(doc, 2, "Preencha os dados da conta", "Informe primeiro nome, sobrenome, nome de usuário, e-mail e uma senha com pelo menos oito caracteres.")
    add_field_table(doc, [
        ("Primeiro nome", "Nome da pessoa responsável pela conta.", "Não use o nome da casa neste campo."),
        ("Sobrenome", "Sobrenome da pessoa responsável.", "Use dados verdadeiros para facilitar a validação."),
        ("Usuário", "Identificador de 3 a 40 caracteres.", "Use letras sem acento, números, ponto, hífen ou underline."),
        ("E-mail", "E-mail individual e acessível.", "É para este e-mail que poderão ser enviados códigos de confirmação."),
        ("Senha", "No mínimo 8 caracteres.", "Crie uma senha exclusiva e não a compartilhe."),
        ("Cidade, bairro e CEP", "Localização-base opcional.", "Se iniciar esse bloco, preencha os três; o CEP precisa ter 8 dígitos."),
    ])
    add_image_placeholder(doc, 2, "Formulário Criar conta", "Formulário completo, do Primeiro nome até CEP, com o botão “Criar conta”.", "Use dados fictícios; não mostre a senha digitada.")
    add_numbered_step(doc, 3, "Finalize o cadastro", "Revise o e-mail e toque em Criar conta. Aguarde a entrada automática no aplicativo. Se aparecer uma mensagem de validação, corrija o campo indicado e tente novamente.")
    add_image_placeholder(doc, 3, "Conta criada e usuário autenticado", "Primeira tela exibida após a criação, mostrando que a pessoa está conectada.")
    add_note(doc, "SE A PESSOA JÁ TEM CONTA", "Não crie outra. Use Já tenho conta e entre com o e-mail e a senha existentes.", BLUE)

    # Venue claim
    doc.add_page_break()
    add_heading(doc, "2  Cadastro da casa e solicitação de acesso", 1)
    add_body(doc, "O fluxo evita cadastros duplicados. Primeiro procure a casa na base do 77Gira. Se ela não aparecer, envie a solicitação interna de inclusão disponível na própria tela. A solicitação recebe protocolo e não publica a casa nem libera acesso automaticamente.")
    add_numbered_step(doc, 1, "Abra Conta e preferências", "Em Configurações, toque no menu de três pontos do cartão da conta. Na seção Perfis e acessos, selecione Administro uma casa.")
    add_image_placeholder(doc, 4, "Acesso a Conta e preferências", "Configurações com o menu de três pontos da conta e, em seguida, a seção “Perfis e acessos”.")
    add_image_placeholder(doc, 5, "Opção Administro uma casa", "Linha “Administro uma casa — Encontre a casa e solicite acesso à gestão”.")
    add_numbered_step(doc, 2, "Pesquise a casa", "Use nome, bairro, região ou cidade. Confira endereço e localização antes de escolher, principalmente quando houver nomes semelhantes.")
    add_image_placeholder(doc, 6, "Busca da casa", "Tela “Encontre a casa que você administra”, com a busca preenchida e o resultado correto visível.")
    # Keep the first field table together with the step that introduces it.
    doc.add_page_break()
    add_numbered_step(doc, 3, "Se a casa já estiver cadastrada", "No resultado correspondente, toque em Solicitar acesso. Confira nome e localização antes de continuar e preencha os dados do vínculo com informações verdadeiras e atuais.")
    add_field_table(doc, [
        ("Tipo de vínculo", "Proprietário ou representante legal, ou Equipe autorizada.", "Escolha a relação que existe de fato."),
        ("Nome do responsável", "Pessoa que responde pela solicitação.", "Mínimo de 3 caracteres."),
        ("Telefone ou WhatsApp", "Contato para confirmação.", "Use um número ativo; mínimo de 8 caracteres."),
        ("CPF ou CNPJ", "Documento do solicitante.", "Revise os números antes de enviar."),
        ("Função ou relação", "Ex.: proprietário, sócio ou gerente.", "Explique de forma objetiva."),
        ("Confirmação do vínculo", "Diga como a equipe pode verificar a autorização.", "Entre 5 e 500 caracteres; inclua referência verificável sem expor senhas."),
    ])
    add_image_placeholder(doc, 7, "Formulário Solicitar acesso", "Modal com o nome da casa e todos os campos da solicitação.", "Oculte CPF/CNPJ, telefone e qualquer dado real.")
    add_numbered_step(doc, 4, "Se a casa não estiver cadastrada", "Vá até o bloco Não encontrou a casa, disponível mesmo quando outros resultados aparecem, e toque em Solicitar inclusão de casa. Não escolha uma casa parecida apenas para prosseguir.")
    add_field_table(doc, [
        ("Nome de exibição", "Nome pelo qual o público encontra a casa.", "Ex.: Tarana Casa de Samba; não use o nome da pessoa responsável."),
        ("Endereço completo", "Rua, número e complemento, quando houver.", "Revise antes de enviar para evitar rotas incorretas."),
        ("Bairro e região", "Localização usada na organização territorial.", "Ex.: Centro, Zona Norte ou a região aplicável."),
        ("Cidade e UF", "Município e estado da casa.", "A UF deve ter duas letras."),
        ("Instagram oficial", "URL completa do perfil oficial.", "Opcional; não informe perfil pessoal."),
        ("Responsável e telefone", "Pessoa e contato para conferência.", "Use dados atuais e acessíveis."),
        ("CPF ou CNPJ", "Documento do solicitante.", "Revise os números e não os exponha em capturas."),
        ("Função ou relação", "Ex.: proprietário, sócio, gerente ou produtor responsável.", "Descreva a relação que existe de fato."),
        ("Confirmação do vínculo", "Como a equipe pode verificar a autorização.", "Inclua referência verificável sem informar senha ou código."),
    ])
    add_image_placeholder(doc, 8, "Inclusão de casa não encontrada", "Faça duas capturas: o bloco persistente “Não encontrou a casa?” e o formulário “Solicitar inclusão de casa” preenchido com dados fictícios.", "Oculte CPF/CNPJ, telefone e qualquer dado real.", 1.8)
    add_numbered_step(doc, 5, "Registre a ciência de responsabilidade", "Ao enviar qualquer uma das solicitações, leia integralmente a Declaração de Responsabilidade e Legitimidade. Role o texto até o final, marque a declaração e toque em Estou ciente e desejo continuar.")
    add_image_placeholder(doc, 9, "Declaração de responsabilidade", "Modal da declaração com o final do texto, checkbox liberado e botão “Estou ciente e desejo continuar”.")
    add_numbered_step(doc, 6, "Aceite os documentos aplicáveis, se a tela aparecer", "Leia cada versão exibida, marque Li e aceito esta versão do documento e toque em Registrar aceite e voltar à ação. Esse aceite não substitui a assinatura formal da etapa seguinte quando ela for exigida.")
    add_image_placeholder(doc, 10, "Aceite documental da solicitação", "Tela “Antes de reivindicar a casa”, documentos, checkboxes e ação final.")
    doc.add_page_break()
    add_numbered_step(doc, 7, "Confirme o envio e guarde o protocolo", "Para uma casa existente, o resultado passa a indicar Em análise. Para uma inclusão, a confirmação exibe um protocolo RA e a seção Inclusões em andamento mostra o nome solicitado. Não envie outro pedido para a mesma casa.")
    add_image_placeholder(doc, 11, "Solicitação em análise", "Faça uma captura do resultado com status “Em análise” ou, no caso de inclusão, da mensagem com protocolo e da seção “Inclusões em andamento”.")

    # Approval and signature
    doc.add_page_break()
    add_heading(doc, "3  Aprovação e assinatura para liberar a gestão", 1)
    add_body(doc, "A análise é feita pela equipe 77Gira. Nenhuma solicitação libera acesso automaticamente. Quando a casa veio de uma solicitação de inclusão, a equipe pode vinculá-la a um cadastro interno já existente ou criar um novo rascunho privado. Nos dois casos, o acesso permanece bloqueado até a conclusão da assinatura formal.")
    add_numbered_step(doc, 1, "Acompanhe a situação", "Volte a Configurações. Um indicador na conta sinaliza assinaturas pendentes; na busca da casa, o status pode mudar para Aguardando assinatura.")
    add_image_placeholder(doc, 12, "Aviso de assinatura pendente", "Configurações com o indicador de pendência e o status “Aguardando assinatura”.")
    add_numbered_step(doc, 2, "Abra a assinatura", "Entre em Conta e preferências, localize Assinaturas formais e toque em Ler e assinar no documento referente à casa.")
    add_image_placeholder(doc, 13, "Lista de Assinaturas formais", "Documento pendente, protocolo, prazo, status e botão “Ler e assinar”.")
    add_numbered_step(doc, 3, "Leia o documento e confirme a leitura", "Leia a versão integral, confira o nome da casa e o protocolo e marque: Li o documento integral e confirmo que estou assinando esta versão de forma consciente.")
    add_numbered_step(doc, 4, "Confirme sua identidade", "Digite a mesma senha usada para entrar no 77Gira e toque em Confirmar minha senha e enviar código. O código só é disparado depois dessa confirmação.")
    doc.add_page_break()
    add_numbered_step(doc, 5, "Digite o código e assine", "Abra o e-mail da conta, copie os seis dígitos, volte ao 77Gira e toque em Assinar documento. Aguarde a mensagem Assinatura registrada com sucesso.")
    add_image_placeholder(doc, 14, "Assinatura em três etapas", "Documento aberto mostrando leitura, senha da conta, envio do código e campo de seis dígitos.", "Nunca capture senha nem código reais.", 1.8)
    add_numbered_step(doc, 6, "Atualize o acesso", "Feche o documento, retorne a Configurações e atualize a tela. O Hub de Gestão deverá exibir Gestão de casas. Se você já administrava outra unidade, a nova casa será acrescentada à mesma conta; ela não substitui as anteriores. Se o acesso não aparecer, entre novamente antes de procurar suporte.")
    add_image_placeholder(doc, 15, "Hub de Gestão liberado", "Configurações com o cartão “Gestão de casas” e a ação “Gerenciar casas”.")
    add_note(doc, "SEGURANÇA", "A equipe 77Gira não precisa conhecer sua senha nem o código recebido por e-mail. Se alguém pedir esses dados, não forneça.", RED)

    # House profile
    doc.add_page_break()
    add_heading(doc, "4  Conferência e atualização dos dados da casa", 1)
    add_numbered_step(doc, 1, "Abra Gestão de casas", "No Hub de Gestão, toque em Gestão de casas. Se houver mais de uma unidade, use o seletor de casa e confirme qual está ativa antes de alterar dados, cadastrar produtores, criar eventos, editar o cardápio ou consultar métricas.")
    add_image_placeholder(doc, 16, "Painel Gestão de Agenda da Casa", "Tela da gestão com o nome da unidade ativa e o menu lateral.")
    add_numbered_step(doc, 2, "Abra Dados da Casa", "No menu lateral, escolha Dados da Casa. Confira a ficha administrativa e toque em Solicitar alteração de dados. Uma casa criada por inclusão pode estar como rascunho e ainda não possuir perfil público disponível.")
    add_numbered_step(doc, 3, "Revise o cadastro público e operacional", "Preencha ou corrija somente informações confirmadas. O nome, a descrição, o endereço, o bairro, a região, a cidade e a UF são essenciais para a apresentação e os filtros.")
    add_field_table(doc, [
        ("Nome da casa", "Nome oficial ou público.", "Evite variações duplicadas do mesmo estabelecimento."),
        ("Artigo e preposição", "Escolha a forma natural de citar a casa.", "Confira a prévia, por exemplo “na Casa X”."),
        ("Apelido", "Forma curta conhecida pelo público.", "Opcional; use apenas se for realmente reconhecida."),
        ("Descrição", "Apresentação objetiva da casa.", "Informe proposta, ambiente e diferenciais sem promessas enganosas."),
        ("Responsável e telefone", "Contato operacional da casa.", "Mantenha os dados atualizados."),
        ("Instagram", "URL completa do perfil oficial.", "Teste o link antes de salvar."),
        ("Endereço", "Logradouro e número.", "Revise para não prejudicar rotas e busca."),
        ("Latitude e longitude", "Coordenadas para Tô na Pista.", "Use coordenadas exatas; não estime."),
        ("Bairro, região, cidade e UF", "Localização usada nos filtros.", "Escolha a região disponível e use UF com 2 letras."),
        ("Imagem", "URL ou upload JPG, PNG ou WebP, até 5 MB.", "Use imagem autorizada, nítida e representativa."),
        ("Dias de funcionamento", "Ex.: Seg, Qua, Sex, Sab.", "Informe os dias habituais; eventos continuam com datas próprias."),
    ])
    add_image_placeholder(doc, 17, "Formulário Dados da Casa", "Formulário de edição com os principais dados e a prévia de tratamento textual.", "Faça duas capturas se o formulário não couber em uma tela.", 1.8)
    add_numbered_step(doc, 4, "Envie a alteração", "Inclua a justificativa obrigatória para a revisão administrativa e toque em Enviar solicitação. Não repita o envio enquanto a solicitação estiver pendente.")
    add_note(doc, "NÃO ALTERE POR TENTATIVA", "Nome, endereço, região e coordenadas afetam busca, rotas e eventos. Em caso de dúvida, confirme os dados antes de enviar.", RED)

    # Producer
    doc.add_page_break()
    add_heading(doc, "5  Cadastro e vínculo de produtor", 1)
    add_body(doc, "Cada produtor deve ter conta própria. A casa pode vincular uma conta existente ou criar uma conta inicial e, em seguida, vinculá-la. A forma recomendada é a pessoa criar a própria conta e informar o e-mail usado.")
    add_heading(doc, "Opção A  Vincular um produtor que já tem conta", 2)
    add_numbered_step(doc, 1, "Abra Produtores", "Na Gestão de casas, confirme a unidade ativa e toque em Produtores.")
    add_numbered_step(doc, 2, "Pesquise a pessoa", "No campo Buscar produtor por nome, email ou usuário, digite uma informação exata. Selecione a pessoa correta na lista, conferindo nome e e-mail.")
    add_numbered_step(doc, 3, "Crie o vínculo", "Toque em Vincular produtor e aguarde a mensagem Produtor vinculado com sucesso. Confira se ele aparece na lista da casa.")
    add_image_placeholder(doc, 18, "Vincular produtor existente", "Área Produtores com a unidade, campo de busca, lista de resultado e botão “Vincular produtor”.", "Oculte o e-mail real na versão pública do manual.")

    add_heading(doc, "Opção B  Criar a conta inicial do produtor", 2)
    add_numbered_step(doc, 1, "Confirme que a pessoa ainda não tem conta", "Evite contas duplicadas. Peça autorização ao produtor e confirme o e-mail individual que será usado.")
    add_numbered_step(doc, 2, "Preencha a criação", "Informe nome, sobrenome, usuário de acesso, e-mail, telefone opcional e uma senha provisória com pelo menos seis caracteres.")
    add_field_table(doc, [
        ("Nome e sobrenome", "Dados pessoais do produtor.", "Não use o nome da casa."),
        ("Usuário de acesso", "Identificador único da pessoa.", "Escolha algo fácil de reconhecer e sem dados desnecessários."),
        ("E-mail", "E-mail individual do produtor.", "É indispensável para acesso e confirmações."),
        ("Telefone", "Contato opcional.", "Cadastre apenas com ciência da pessoa."),
        ("Senha provisória", "Mínimo de 6 caracteres na tela atual.", "Envie por canal privado e peça alteração imediata para uma senha exclusiva de 8 ou mais caracteres."),
    ])
    add_numbered_step(doc, 3, "Crie e depois vincule", "Toque em Criar produtor. A criação da conta, sozinha, não conclui o processo: pesquise o produtor no formulário de vínculo, selecione-o e toque em Vincular produtor.")
    add_image_placeholder(doc, 19, "Criar produtor e concluir o vínculo", "Formulário de criação e, abaixo, formulário de vínculo com o produtor recém-criado selecionado.", "Nunca deixe a senha provisória visível na captura.", 1.8)
    add_numbered_step(doc, 4, "Peça a troca da senha provisória", "O produtor entra na conta, abre Configurações, toca no menu de três pontos, escolhe editar os dados pessoais e usa Alterar senha. Ao concluir, as demais sessões abertas são encerradas.")
    add_heading(doc, "Opção C  Solicitação feita pelo próprio produtor", 2)
    add_body(doc, "Se o produtor ainda não está vinculado pela casa, ele pode criar a própria conta e escolher Conta e preferências → Perfis e acessos → Atuo como produtor. Ele pesquisa a casa, solicita acesso, registra as ciências exigidas, aguarda a análise e assina o documento quando solicitado.")
    add_note(doc, "REGRA DE ACESSO", "Remover um vínculo retira o acesso do produtor àquela casa. Não compartilhe a conta do gestor como atalho.", RED)

    # Event
    doc.add_page_break()
    add_heading(doc, "6  Cadastro e publicação de evento", 1)
    add_numbered_step(doc, 1, "Abra Eventos", "Na Gestão de casas, confirme a unidade ativa e escolha Eventos. O formulário mostra uma prévia do card.")
    add_image_placeholder(doc, 20, "Entrada da programação", "Menu da casa com Eventos selecionado e a prévia do card visível.")
    add_numbered_step(doc, 2, "Preencha as informações principais", "Use informações confirmadas pela produção. Data e hora de término precisam ser posteriores ao início.")
    add_field_table(doc, [
        ("Título", "Nome do evento ou artista, se for show solo.", "Obrigatório; deixe claro para o público."),
        ("Artista principal", "Escolha ou digite o artista.", "Opcional; confira a grafia oficial."),
        ("Descrição", "Resumo do evento.", "Obrigatória; informe formato, atrações e condições relevantes."),
        ("Tipo", "Roda de Samba, Pagode, Gafieira, Samba Rock ou Samba com Feijoada.", "Escolha o gênero predominante."),
        ("Casa", "Unidade onde ocorre o evento.", "Para o perfil da casa, vem selecionada e bloqueada."),
        ("Início e fim", "Data e hora completas.", "O fim deve ser depois do início."),
        ("Ingresso", "Pago, Gratuito ou Consumação.", "Em Gratuito, não informe preço mínimo ou máximo."),
        ("Valores", "Preço mínimo, máximo, consumação e couvert quando aplicáveis.", "Use valores numéricos e revise a política do evento."),
        ("URL de ingresso", "Link oficial de compra ou reserva.", "Opcional; teste antes de publicar."),
        ("Cartaz", "URL ou upload JPG, PNG ou WebP, até 5 MB.", "Recomendado: 1080 × 1350 px, proporção 4:5."),
        ("Sinalizações", "Samba Familiar e Kids Friendly, quando verdadeiras.", "Não marque apenas para ampliar alcance."),
        ("Outras tags", "Termos separados por vírgula.", "Opcional; evite repetições."),
    ])
    add_image_placeholder(doc, 21, "Formulário de evento", "Formulário completo do título às tags e os botões “Criar rascunho” e “Publicar evento”.", "Use duas capturas verticais se necessário.", 1.8)
    add_heading(doc, "Evento recorrente semanal", 2)
    add_numbered_step(doc, 3, "Ative somente quando a programação se repetir", "Marque Evento recorrente semanal, selecione os dias, informe os horários da recorrência, a data final e, se já souber, as datas sem evento no formato AAAA-MM-DD separadas por vírgula.")
    add_image_placeholder(doc, 22, "Configuração de recorrência", "Painel com dias da semana, horários, data final e exceções.")
    add_heading(doc, "Salvar ou publicar", 2)
    add_numbered_step(doc, 4, "Use rascunho quando ainda faltar confirmação", "Toque em Criar rascunho. O evento aparece na lista como Rascunho e não deve ser considerado publicado para o público.")
    add_numbered_step(doc, 5, "Faça a revisão obrigatória", "Para colocar no ar, toque em Publicar evento. Confira o resumo e marque os seis itens: título/artista, data e horário, casa, região, preço e mídia/descrição/link.")
    add_numbered_step(doc, 6, "Publique", "Quando todos os itens estiverem marcados, toque em Publicar agora. Aguarde Evento publicado com sucesso e confirme o status Publicado na lista.")
    add_image_placeholder(doc, 23, "Revisão obrigatória antes de publicar", "Modal com resumo, seis itens de verificação e botão “Publicar agora”.")
    add_image_placeholder(doc, 24, "Evento publicado na lista", "Card do evento com status “Publicado” e ações de gerenciamento.")
    add_heading(doc, "Depois da publicação", 2)
    add_bullet(doc, "Use Editar para corrigir os dados e repita a revisão de publicação quando necessário.")
    add_bullet(doc, "Em séries recorrentes, use Cancelar data para uma ocorrência específica e Reativar data para restaurá-la.")
    add_bullet(doc, "Use Excluir somente quando a remoção for realmente necessária; prefira corrigir ou cancelar a ocorrência quando aplicável.")
    add_note(doc, "CHECAGEM PÚBLICA", "Depois de publicar, abra o Explorar e confirme título, casa, região, horário, preço, imagem e link. Não presuma que o cadastro correto garante uma apresentação perfeita sem conferência.", GREEN)

    # Menu
    doc.add_page_break()
    add_heading(doc, "7  Cadastro e publicação do Cardápio Essencial", 1)
    add_body(doc, "O Cardápio Essencial é uma seleção curta e fácil de manter. A tela aceita até 30 itens ativos e permite ocultar preços, arquivar itens e importar uma planilha CSV.")
    add_numbered_step(doc, 1, "Abra o cardápio", "Na Gestão de casas, confirme a unidade ativa e escolha Cardápio Essencial. Se a tela mostrar as condições do modelo gratuito, leia antes de prosseguir.")
    add_image_placeholder(doc, 25, "Tela inicial do Cardápio Essencial", "Cabeçalho, aviso do modelo gratuito, status, contagem de itens e controle de preços.")
    add_numbered_step(doc, 2, "Confirme a condição apresentada", "Marque Li e aceito as condições de publicidade do Cardápio Essencial e toque em Confirmar condição. O botão de publicação só fica disponível depois desse registro na interface atual.")
    add_note(doc, "LEIA ANTES DE ACEITAR", "A tela informa que a publicidade, quando exibida, é selecionada e distribuída pelo 77Gira; a casa não escolhe anunciantes nem recebe participação por essas exibições.", BLUE)
    add_numbered_step(doc, 3, "Defina se os preços serão exibidos", "No controle Exibir preços, mantenha Sim para mostrar valores. Desative apenas quando a opção comercial da casa exigir consulta ou quando os itens estiverem configurados adequadamente.")
    add_numbered_step(doc, 4, "Cadastre o primeiro item", "Preencha o bloco Novo item e toque em Adicionar item. Repita para cada item da seleção básica.")
    add_field_table(doc, [
        ("Nome", "Nome do prato ou bebida.", "Obrigatório; máximo de 100 caracteres."),
        ("Categoria", "Petiscos, Porções, Pratos, Lanches, Sobremesas, Cervejas, Drinks, Doses, Vinhos e espumantes ou Sem álcool.", "Escolha a categoria que facilita a leitura."),
        ("Descrição", "Ingredientes, preparo ou composição.", "Opcional; máximo de 240 caracteres."),
        ("Preço", "Valor em reais, por exemplo 29,90.", "Deixe vazio somente quando a modalidade permitir."),
        ("Modalidade", "Preço exato, A partir de, Oculto ou Sob consulta.", "Combine com a forma real de venda."),
        ("Apresentação", "Individual, serve 2, unidade, dose, garrafa, lata, porção, jarra, balde etc.", "Opcional; escolha a unidade correta."),
        ("Status", "Disponível, Indisponível ou Rascunho.", "Use Rascunho para item incompleto."),
        ("Características", "Até 4 marcadores, como vegetariano ou especialidade da casa.", "Marque somente características verdadeiras."),
        ("Destaque", "Destacar este item no cardápio.", "Reserve para poucos itens importantes."),
    ])
    add_image_placeholder(doc, 26, "Formulário Novo item", "Todos os campos do item e o botão “Adicionar item”.")
    add_numbered_step(doc, 5, "Revise a lista", "Confirme nome, categoria, status e destaque. Use as setas para ordenar, Editar para corrigir e Arquivar para retirar um item sem apagá-lo definitivamente.")
    add_image_placeholder(doc, 27, "Lista de itens do cardápio", "Abas Ativos e Arquivados, ordem dos itens e ações Editar e Arquivar.")
    add_numbered_step(doc, 6, "Publique o cardápio", "Quando a seleção estiver correta, toque em Publicar e marcar revisado. Aguarde a mensagem Cardápio publicado e marcado como revisado e confirme o status Publicado.")
    add_image_placeholder(doc, 28, "Cardápio publicado", "Barra superior com status “Publicado”, contagem de itens e confirmação da publicação.")
    add_heading(doc, "Atualizações do cardápio", 2)
    add_bullet(doc, "Item temporariamente indisponível: edite e altere o status para Indisponível.")
    add_bullet(doc, "Item que saiu da seleção: use Arquivar. Ele permanecerá na aba Arquivados.")
    add_bullet(doc, "Item arquivado que voltou: use Restaurar como rascunho, revise e publique novamente.")
    add_bullet(doc, "Mudança de ordem: use as setas para cima e para baixo e confira o resultado.")
    add_heading(doc, "Importação por planilha CSV  Opcional", 2)
    add_numbered_step(doc, 1, "Baixe o modelo", "Na seção Planilha do cardápio, toque em Baixar modelo CSV. Edite em formato CSV UTF-8.")
    add_numbered_step(doc, 2, "Selecione e valide", "Toque em Selecionar CSV. O arquivo pode ter até 512 KB; a prévia informa erros antes de salvar. A importação não substitui itens existentes e o total não pode passar de 30 itens ativos.")
    add_numbered_step(doc, 3, "Confirme", "Revise a quantidade e toque em Confirmar importação. Os itens entram como um único lote; depois, revise e publique o cardápio.")
    add_image_placeholder(doc, 29, "Ferramentas de planilha", "Botões de baixar modelo, baixar cardápio, selecionar CSV e a prévia de validação.")

    # Final checks and troubleshooting
    doc.add_page_break()
    add_heading(doc, "8  Verificação final antes de divulgar", 1)
    add_checklist(doc, [
        "A conta pertence a uma pessoa real e usa e-mail individual.",
        "A casa correta está vinculada; não há cadastro duplicado e o protocolo da inclusão foi guardado quando aplicável.",
        "A solicitação foi aprovada e a assinatura formal, quando exigida, está concluída.",
        "Nome, descrição, endereço, bairro, região, cidade, UF, dias e imagem da casa foram revisados.",
        "Cada produtor tem conta própria e aparece na lista de produtores da unidade correta.",
        "O primeiro evento está com status Publicado e foi conferido no Explorar.",
        "O cardápio tem itens corretos, preços coerentes e status Publicado.",
        "Links de ingresso e Instagram foram testados.",
        "Nenhuma senha, código ou documento pessoal foi compartilhado em grupo ou captura de tela.",
    ])
    add_image_placeholder(doc, 30, "Perfil público e agenda conferidos", "Visão pública da casa e o evento no Explorar, sem informações administrativas.")
    add_image_placeholder(doc, 31, "Cardápio visto pelo público", "Cardápio publicado na visão pública, com itens, preços e ordem final.")

    doc.add_page_break()
    add_heading(doc, "Se algo não funcionar", 1)
    issues = [
        ("Não encontro a casa", "Refaça a busca por nome, bairro e cidade. Se não existir, use o bloco Não encontrou a casa e toque em Solicitar inclusão de casa; não duplique."),
        ("A inclusão continua Em análise", "Confira Inclusões em andamento e guarde o protocolo RA. Aguarde a equipe vincular o pedido a uma casa existente ou criar o rascunho; não envie solicitações repetidas."),
        ("Aparece Aguardando assinatura", "Abra Configurações → Conta e preferências → Assinaturas formais e conclua as três etapas."),
        ("O código não chegou", "Confirme primeiro a senha da conta, verifique spam e se o e-mail exibido pertence à conta. Use Reenviar código quando disponível."),
        ("Gestão de casas não aparece", "Atualize a tela e entre novamente. Confirme se a assinatura ficou com status Assinado. Se você administra várias casas, confira o seletor de unidade."),
        ("Não encontro o produtor", "Confirme se ele já tem conta e pesquise por e-mail exato ou usuário. Se acabou de ser criado, atualize a busca antes de vincular."),
        ("Evento não publica", "Revise campos destacados, término depois do início, política de preço e todos os seis itens da revisão obrigatória."),
        ("Cardápio não publica", "Confirme a condição exibida, revise os itens e verifique se o botão Publicar e marcar revisado foi liberado."),
        ("Imagem não envia", "Use JPG, PNG ou WebP com até 5 MB. Para cartaz de evento, prefira 1080 × 1350 px."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    issue_width = Inches(2.45)
    action_width = Inches(4.3)
    table.columns[0].width = issue_width
    table.columns[1].width = action_width
    for i, text in enumerate(("Situação", "O que fazer")):
        set_cell_width(table.rows[0].cells[i], (issue_width, action_width)[i])
        set_cell_shading(table.rows[0].cells[i], NAVY)
        add_run(table.rows[0].cells[i].paragraphs[0], text, bold=True, color="FFFFFF", size=9)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for situation, action in issues:
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        set_cell_width(cells[0], issue_width)
        set_cell_width(cells[1], action_width)
        add_run(cells[0].paragraphs[0], situation, bold=True, color=NAVY, size=8.6)
        add_run(cells[1].paragraphs[0], action, size=8.6)
        for cell in cells:
            border = {"val": "single", "sz": "4", "color": LIGHT}
            set_cell_border(cell, top=border, bottom=border, left=border, right=border)

    add_heading(doc, "Boas práticas permanentes", 1)
    add_bullet(doc, "Mantenha os dados da casa e a agenda atualizados.")
    add_bullet(doc, "Publique apenas informações confirmadas e imagens com autorização de uso.")
    add_bullet(doc, "Retire prontamente eventos cancelados e itens indisponíveis.")
    add_bullet(doc, "Revise acessos da equipe quando alguém deixar a operação da casa.")
    add_bullet(doc, "Use canais oficiais do 77Gira para solicitar ajuda e nunca envie senhas ou códigos.")
    add_body(doc, "Fim do manual", before=18, after=0)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()

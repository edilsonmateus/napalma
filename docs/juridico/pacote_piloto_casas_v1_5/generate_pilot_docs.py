from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor


OUTPUT = Path(__file__).parent


def set_font(run, name="Arial", size=11, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in [("Title", 19), ("Heading 1", 14), ("Heading 2", 12)]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(16 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("77Gira | Versão piloto 1.5 | Documento sujeito à revisão jurídica antes da publicação definitiva")
    set_font(run, size=8, color="555555")


def add_paragraph(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_font(run, bold=True)
        run = paragraph.add_run(text[len(bold_prefix):])
        set_font(run)
    else:
        run = paragraph.add_run(text)
        set_font(run)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)


def add_title(doc, title, subtitle):
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(title)
    set_font(run, size=19, bold=True)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(subtitle)
    set_font(run, size=10, color="555555")
    paragraph.paragraph_format.space_after = Pt(18)


def add_section(doc, title, paragraphs=(), bullets=()):
    heading = doc.add_paragraph(style="Heading 1")
    run = heading.add_run(title)
    set_font(run, size=14, bold=True)
    for paragraph in paragraphs:
        add_paragraph(doc, paragraph)
    if bullets:
        add_bullets(doc, bullets)


def acceptance_text(title, subtitle, sections):
    lines = [title, subtitle, ""]
    for section in sections:
        lines.append(section["title"])
        lines.extend(section.get("paragraphs", []))
        lines.extend(f"- {item}" for item in section.get("bullets", []))
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def save_document(filename, title, subtitle, sections):
    doc = Document()
    configure(doc)
    add_title(doc, title, subtitle)
    for section in sections:
        add_section(doc, **section)
    destination = OUTPUT / filename
    doc.save(destination)
    destination.with_suffix(".txt").write_text(acceptance_text(title, subtitle, sections), encoding="utf-8")


COMMON_IDENTIFICATION = (
    "77Gira é uma iniciativa digital operada por Edilson Mateus de Oliveira, com canal de contato em "
    "77giramundo@gmail.com e endereço para correspondência na Rua Antonio Martins Costa, 407, Jardim Boa Vista, "
    "Butantã, São Paulo, SP, CEP 05584-000."
)


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)

    save_document(
        "01_termos_de_uso_77gira_piloto_v1_5.docx",
        "Termos de Uso 77Gira Piloto",
        "Versão 1.5 | Aplicável à etapa inicial de operação com casas e programação",
        [
            {"title": "1. Objeto e etapa piloto", "paragraphs": [
                "Estes Termos disciplinam o uso do 77Gira na etapa inicial de operação com casas, produtores, artistas e público. A etapa piloto serve para organizar perfis, divulgar programação e testar os fluxos de gestão em escala controlada.",
                "A participação da casa no piloto é gratuita. Nesta etapa, não há contratação de mídia, cobrança, venda de créditos, patrocínio, publicidade paga ou promessa de retorno comercial. Qualquer operação comercial futura dependerá de instrumento próprio e de aceite específico."
            ]},
            {"title": "2. Conta e representação", "paragraphs": [
                COMMON_IDENTIFICATION,
                "Quem usa o 77Gira em nome de uma casa, artista, produtor ou organização declara possuir autorização legítima para fazê-lo. A pessoa deve manter seus dados corretos, proteger suas credenciais e conceder acesso apenas a integrantes autorizados da operação."
            ]},
            {"title": "3. Casas, produtores e programação", "paragraphs": [
                "A casa ou pessoa autorizada responde pela exatidão, atualização e legitimidade dos dados publicados, incluindo nome de exibição, endereço, contatos, imagens, horário de funcionamento, programação, preços, classificação e avisos relevantes ao público.",
                "A pessoa que solicita ou recebe gestão de uma casa responde pela escolha dos produtores vinculados e pelas permissões concedidas. O 77Gira pode pedir confirmação, corrigir, limitar, pausar ou remover conteúdos e acessos quando houver inconsistência, risco, fraude, violação de direito ou necessidade operacional."
            ]},
            {"title": "4. Conteúdo e direitos", "paragraphs": [
                "Quem envia imagem, marca, texto, agenda ou outro material declara possuir os direitos e autorizações necessários. Para operar o serviço, a pessoa concede ao 77Gira licença não exclusiva, gratuita e limitada para armazenar, adaptar tecnicamente e exibir o conteúdo nos canais da plataforma enquanto ele permanecer publicado.",
                "A programação é informativa e pode mudar. O público deve confirmar diretamente com a casa ou produção detalhes como horário, valor, lotação, classificação e realização do evento."
            ]},
            {"title": "5. Reivindicação e aceite", "paragraphs": [
                "Antes de enviar uma reivindicação ou pedido de acesso profissional, o solicitante deverá ler e aceitar o Termo de Reivindicação e Gestão de Perfil no Piloto. O registro vincula conta, versão do documento, data, horário e evidências técnicas disponíveis.",
                "A análise da reivindicação não garante aprovação. Quando a solicitação for aprovada, a plataforma poderá pedir assinatura eletrônica reforçada do mesmo termo antes de liberar o acesso de gestão."
            ]},
            {"title": "6. Publicidade e negociação comercial", "paragraphs": [
                "A 77Gira Ads não integra este piloto. Não haverá publicidade paga, cobrança, patrocínio, compra de mídia, carteira de créditos ou compartilhamento de dados de público com anunciantes nesta etapa.",
                "Se a plataforma iniciar operação comercial posteriormente, as condições aplicáveis serão apresentadas de forma separada e não alterarão retroativamente a participação gratuita no piloto."
            ]},
            {"title": "7. Condutas vedadas e medidas", "bullets": [
                "Assumir perfil, publicar ou alterar dados de terceiros sem autorização.",
                "Inserir informações enganosas, ilegais, discriminatórias, fraudulentas ou que violem direitos de terceiros.",
                "Compartilhar credenciais, burlar permissões, explorar falhas ou extrair dados em massa sem autorização.",
                "Usar o serviço para spam, assédio, fraude ou atividade incompatível com a finalidade cultural e informativa da plataforma."
            ]},
            {"title": "8. Disponibilidade, encerramento e contato", "paragraphs": [
                "Funcionalidades do piloto podem ser ajustadas, interrompidas, pausadas ou encerradas para manutenção, segurança, melhoria do produto ou decisão operacional. Isso não elimina os registros que precisem ser conservados para segurança, auditoria ou cumprimento de obrigação legal.",
                "Dúvidas sobre estes Termos, conteúdo ou participação no piloto podem ser encaminhadas para 77giramundo@gmail.com. A Política de Privacidade complementa estes Termos quanto ao tratamento de dados pessoais."
            ]},
        ],
    )

    save_document(
        "02_politica_de_privacidade_77gira_piloto_v1_5.docx",
        "Política de Privacidade 77Gira Piloto",
        "Versão 1.5 | Transparência sobre dados pessoais na etapa inicial com casas e programação",
        [
            {"title": "1. Controlador e alcance", "paragraphs": [
                COMMON_IDENTIFICATION,
                "Nesta etapa, Edilson Mateus de Oliveira atua como controlador dos dados pessoais tratados pelo 77Gira e define as finalidades essenciais do tratamento. Esta política se aplica às páginas públicas, contas, reivindicações, perfis de casas, programação, suporte e áreas de gestão."
            ]},
            {"title": "2. Dados tratados no piloto", "bullets": [
                "Dados de conta e contato: nome, e-mail, telefone, nome de usuário, senha protegida por hash e dados necessários à autenticação.",
                "Dados profissionais e de representação: vínculo declarado, função, evidências de legitimidade, decisões de reivindicação, permissões, assinatura e registros de auditoria.",
                "Dados de casas e programação: nome, endereço comercial, contatos divulgados, imagens, agenda, preços, horários e demais informações publicadas por pessoas autorizadas.",
                "Dados técnicos e de segurança: registros de acesso, sessão, navegador, identificadores técnicos, sinais de fraude e logs necessários para proteger a plataforma.",
                "Pedidos de suporte e de privacidade: conteúdo da solicitação, confirmação de identidade, protocolo e histórico de atendimento."
            ]},
            {"title": "3. Finalidades e bases legais", "paragraphs": [
                "Usamos os dados para criar e proteger contas, analisar reivindicações, liberar acessos compatíveis, exibir programação pública, atender suporte, prevenir fraude, manter auditoria e cumprir obrigações legais. Conforme o caso, o tratamento pode se apoiar em procedimentos preliminares e execução da relação solicitada pelo titular, cumprimento de obrigação legal, exercício regular de direitos, legítimo interesse com salvaguardas adequadas ou consentimento quando este for a base apropriada.",
                "Quem fornece dados de outra pessoa deve ter autorização ou outra base legítima aplicável e limitar o envio ao necessário."
            ]},
            {"title": "4. Dados públicos e compartilhamento", "paragraphs": [
                "Dados publicados em página de casa ou evento podem ser acessados por pessoas sem conta. A pessoa autorizada deve decidir cuidadosamente o que será tornado público e manter essas informações atualizadas.",
                "O 77Gira pode utilizar operadores de hospedagem, banco de dados, armazenamento, segurança, comunicação e suporte estritamente para operar a plataforma. O acesso é limitado ao necessário e sujeito às obrigações aplicáveis. Não vendemos dados pessoais."
            ]},
            {"title": "5. O que não faz parte deste piloto", "paragraphs": [
                "A 77Gira Ads não está em operação nesta etapa. Não há publicidade regional, segmentação publicitária, carteira de mídia, cobrança, patrocínio pago nem compartilhamento de dados individuais de usuários com anunciantes.",
                "Recursos opcionais que dependam de localização atual ou notificações permanecem sujeitos às permissões específicas do navegador ou dispositivo."
            ]},
            {"title": "6. Segurança, retenção e direitos", "paragraphs": [
                "O 77Gira adota medidas proporcionais de segurança, incluindo proteção de credenciais, controles de acesso, validação de dados, limitação de requisições e trilhas de auditoria. Nenhum sistema é invulnerável; por isso, não compartilhe senha ou códigos de confirmação.",
                "Conservamos os dados pelo tempo necessário para a finalidade, segurança, auditoria, defesa de direitos e obrigações aplicáveis. O titular pode solicitar confirmação de tratamento, acesso, correção, informações sobre compartilhamento, eliminação ou revogação de consentimento quando cabível, além de outros direitos previstos na LGPD.",
                "Solicitações podem ser feitas pela Central de Privacidade e Dados, quando disponível na conta, ou por 77giramundo@gmail.com. Poderemos confirmar a identidade antes de atender o pedido."
            ]},
            {"title": "7. Atualizações", "paragraphs": [
                "Esta política poderá ser atualizada conforme a evolução do piloto, a entrada de novos recursos ou mudanças legais. Mudanças materiais serão apresentadas por meio adequado e, quando exigirem consentimento, dependerão de nova escolha do titular."
            ]},
        ],
    )

    save_document(
        "03_termo_de_reivindicacao_e_gestao_de_perfil_piloto_v1_5.docx",
        "Termo de Reivindicação e Gestão de Perfil no Piloto",
        "Versão 1.5 | Aplicável a pedidos de gestão de casas, programação e perfis profissionais",
        [
            {"title": "1. Declaração de legitimidade", "paragraphs": [
                "Ao aceitar este Termo, a pessoa solicitante declara, sob sua responsabilidade, que é proprietária, integrante autorizada, representante legal, gestora ou profissional expressamente autorizado a solicitar a criação, reivindicação, administração ou alteração do perfil indicado.",
                "A declaração não substitui a análise do 77Gira e não garante a aprovação do pedido."
            ]},
            {"title": "2. Finalidade do piloto", "paragraphs": [
                "O acesso solicitado destina-se exclusivamente à operação piloto gratuita do 77Gira: organizar o perfil da casa, vincular produtores autorizados, manter a programação e disponibilizar informações corretas ao público.",
                "Este Termo não cria contratação de publicidade, patrocínio, compra de mídia, exclusividade, promessa de receita ou obrigação comercial futura."
            ]},
            {"title": "3. Responsabilidades de gestão", "bullets": [
                "Manter corretos e atualizados os dados da casa, seus contatos, endereço, imagens, programação, preços, horários e avisos relevantes.",
                "Vincular somente produtores e integrantes autorizados e revisar periodicamente os acessos concedidos.",
                "Garantir autorização para textos, fotos, marcas, vídeos, contatos e outros conteúdos publicados.",
                "Comunicar mudanças relevantes, cancelamentos e inconsistências que possam afetar o público.",
                "Não utilizar o acesso para assumir perfis, desviar contatos, publicar informação enganosa ou agir em nome de terceiros sem autorização."
            ]},
            {"title": "4. Verificação e medidas", "paragraphs": [
                "O 77Gira pode solicitar informações, documentos e esclarecimentos antes ou depois de decidir. Pode também pausar publicação, limitar permissões, remover vínculos ou cancelar a solicitação quando houver indício de fraude, falsa representação, risco ao público ou violação deste Termo.",
                "A pessoa solicitante autoriza o registro de protocolo, versão, hash de integridade, conta autenticada, data, hora, evidências e decisões necessárias à auditoria do fluxo."
            ]},
            {"title": "5. Aceite e assinatura", "paragraphs": [
                "O aceite inicial é eletrônico e fica vinculado à conta e à versão deste Termo. Quando a plataforma solicitar assinatura reforçada após a aprovação de elegibilidade, o acesso de gestão será liberado somente depois de concluídas as etapas de confirmação apresentadas ao solicitante.",
                "A assinatura reforçada confirma a mesma versão congelada do documento. Se o Termo for atualizado materialmente antes da assinatura, será apresentada nova versão para leitura e aceite."
            ]},
            {"title": "6. Dados pessoais e contato", "paragraphs": [
                "O tratamento de dados relacionados a este pedido segue a Política de Privacidade do 77Gira. Dúvidas, contestação de decisão ou comunicação de irregularidade podem ser encaminhadas para 77giramundo@gmail.com."
            ]},
        ],
    )

    save_document(
        "04_politica_de_conteudo_moderacao_e_denuncias_piloto_v1_5.docx",
        "Política de Conteúdo Moderação e Denúncias no Piloto",
        "Versão 1.5 | Regras aplicáveis a perfis, agendas e materiais publicados",
        [
            {"title": "1. Conteúdo permitido", "paragraphs": [
                "O 77Gira recebe e exibe informações culturais e operacionais relacionadas a casas, eventos, artistas, produtores e público. As informações publicadas devem ser verdadeiras, atuais, relevantes e acompanhadas das autorizações necessárias."
            ]},
            {"title": "2. Conteúdo proibido", "bullets": [
                "Falsa identidade, reivindicação indevida de perfil, fraude, phishing, spam ou tentativa de obter acesso não autorizado.",
                "Informação enganosa sobre eventos, preços, horários, disponibilidade, atrações ou vínculos profissionais.",
                "Violação de direitos autorais, de imagem, de marca, de privacidade ou de outros direitos de terceiros.",
                "Ameaça, assédio, discriminação, violência, discurso de ódio, conteúdo ilegal ou exposição indevida de dados pessoais.",
                "Links maliciosos, arquivos perigosos, automação abusiva, manipulação de métricas ou exploração de falhas."
            ]},
            {"title": "3. Revisão e medidas", "paragraphs": [
                "O 77Gira pode revisar, pedir correção, limitar alcance, pausar publicação, remover conteúdo, revogar acesso ou suspender conta quando isso for necessário para segurança, proteção de direitos, prevenção de fraude ou cumprimento de obrigação legal.",
                "Sempre que possível, a pessoa responsável receberá orientação para correção. Medidas urgentes podem ser tomadas imediatamente em caso de risco, fraude, conteúdo malicioso, violência ou determinação legal."
            ]},
            {"title": "4. Denúncias e correções", "paragraphs": [
                "Denúncias, pedidos de correção e comunicações sobre uso indevido podem ser enviados para 77giramundo@gmail.com com a identificação da página, do evento ou do conteúdo envolvido e uma descrição objetiva do problema. Nunca envie senha, código de confirmação ou dados pessoais desnecessários.",
                "A plataforma poderá preservar registros e evidências estritamente necessários à análise, à segurança, à resposta ao pedido e ao exercício de direitos."
            ]},
            {"title": "5. Relação com a operação piloto", "paragraphs": [
                "Esta política integra a operação piloto gratuita com casas e programação. Ela não institui publicidade paga, patrocínio, compra de mídia ou relação comercial entre a casa e terceiros."
            ]},
        ],
    )


if __name__ == "__main__":
    main()

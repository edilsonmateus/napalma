from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / "juridico" / "pacote_contratual_v1_4"
OUT.mkdir(parents=True, exist_ok=True)

IDENTITY = "77Gira | 77Giramundo | Edilson Mateus de Oliveira | 77giramundo@gmail.com"
ADDRESS = "Rua Antonio Martins Costa, 407, Jardim Boa Vista - Butantã, São Paulo - SP, CEP 05584-000."
PENDING_J = "[DECISÃO JURÍDICA/PRODUTO PENDENTE]"
PENDING_F = "[DECISÃO FINANCEIRA/JURÍDICA PENDENTE]"

DOCS = [
 ("02_termos_de_uso_77gira_v1_4.docx", "Termos de Uso - 77Gira", "Versão 1.4 | Vigência proposta: 30/08/2026", [
  ("1. Aceite, plataforma e abrangência", ["O 77Gira é uma plataforma da iniciativa 77Giramundo, operada por Edilson Mateus de Oliveira. A sede é São Paulo - SP e a implantação inicial ocorre em São Paulo; a plataforma é concebida para todo o território nacional, com disponibilidade que pode variar conforme a implantação regional.", "Ao criar conta ou aceitar documento apresentado no 77Gira, o Usuário realiza aceite eletrônico vinculado à conta, à versão do documento e aos registros técnicos disponíveis. Quem atua em nome de Casa, Artista, Produtor, Anunciante, marca ou organização declara possuir legitimidade para representá-la."]),
  ("2. Conta, conteúdo e perfis", ["O Usuário deve informar dados verdadeiros, manter a conta protegida e responder por atividades realizadas com suas credenciais. Quem publica conteúdo, imagem, marca, link ou informação declara possuir as autorizações necessárias.", "Reivindicações e acessos profissionais podem ser recusados, revisados ou revogados quando não houver legitimidade suficiente, houver fraude ou violação de direitos."]),
  ("3. Publicidade e serviços de terceiros", ["O 77Gira pode exibir conteúdo identificado como publicitário ou patrocinado. Anunciantes respondem por seus criativos, ofertas, marcas, direitos e destinos. Veiculação depende de aprovação, saldo, período, inventário e regras do posicionamento, sem garantia de alcance, cliques, conversão, vendas, posição, exclusividade ou retorno, salvo compromisso escrito específico.", "Informações de eventos, negociações, mapas e links de terceiros podem mudar. O 77Gira não é automaticamente parte de contratação, venda, organização ou execução de evento."]),
  ("4. Moderação, alterações e contato", ["O 77Gira pode solicitar ajustes, limitar, suspender ou remover conteúdo, acesso ou campanha quando necessário para segurança, cumprimento legal, prevenção de fraude ou proteção de direitos. Mudanças relevantes serão comunicadas por meios adequados.", f"Contato jurídico: 77giramundo@gmail.com. Endereço: {ADDRESS}"])
 ]),
 ("03_politica_de_privacidade_e_cookies_v1_4.docx", "Política de Privacidade e Cookies - 77Gira", "Versão 1.4 | Vigência proposta: 30/08/2026", [
  ("1. Dados e finalidades", ["O 77Gira pode tratar dados de cadastro, conta, conteúdo, vínculos profissionais, reivindicações, eventos, contratação, segurança, publicidade e solicitações de direitos para operar a plataforma, proteger pessoas e contas e cumprir obrigações.", "O tratamento observa a LGPD e a finalidade aplicável. A pessoa pode usar a Central de Privacidade e Dados para consultar preferências e exercer direitos disponíveis."]),
  ("2. Localização e publicidade regional", ["Cidade, bairro e CEP são referências gerais da conta. Localização atual, quando solicitada pelo navegador ou dispositivo, é uma permissão técnica própria e usada somente em recurso que dela dependa.", "Publicidade por região usa somente a cidade-base após decisão afirmativa específica. A recusa não bloqueia o aplicativo e a pessoa ainda pode receber anúncios gerais ou contextuais aprovados. Anunciantes não recebem localização individual, nome, e-mail, telefone, histórico individual ou identidade de quem visualizou ou clicou em publicidade."]),
  ("3. Métricas, cookies e compartilhamento", ["O 77Gira pode registrar entrega, impressão válida, clique, controles de frequência e sinais técnicos de integridade para funcionamento, prevenção de fraude, medição e relatórios agregados. Armazenamento local e de sessão apoia autenticação, preferências, onboarding e funcionamento.", "Dados podem ser tratados por operadores de infraestrutura, comunicação, segurança e armazenamento conforme necessidade. Lista de fornecedores, transferências internacionais e prazos por categoria: " + PENDING_J]),
  ("4. Retenção, segurança e contato", ["Dados são mantidos pelo tempo necessário às finalidades, obrigações, segurança, auditoria e exercício de direitos, com eliminação ou anonimização quando aplicável. A plataforma adota medidas proporcionais de segurança e responderá incidentes conforme a lei.", f"Canal: 77giramundo@gmail.com. Endereço: {ADDRESS}"])
 ]),
 ("04_termos_de_publicidade_77gira_ads_v1_4.docx", "Termos de Publicidade - 77Gira Ads", "Versão 1.4 | Vigência proposta: 30/08/2026", [
  ("1. Conta, campanha e revisão", ["A Conta de Anunciante, campanhas, criativos e destinos podem depender de aprovação. O 77Gira pode aprovar, pedir ajustes, rejeitar, pausar ou encerrar itens por segurança, qualidade, direito de terceiros, fraude, ilegalidade, destino inadequado, inventário ou regra aplicável."]),
  ("2. Responsabilidade do Anunciante", ["O Anunciante responde pelas informações, alegações, ofertas, imagens, marcas, direitos autorais, autorizações de imagem, legitimidade comercial e páginas de destino que fornecer. Aprovação interna não transfere essa responsabilidade ao 77Gira."]),
  ("3. Posicionamentos, entrega e métricas", ["O 77Gira pode criar, alterar, incluir ou descontinuar Posicionamentos Publicitários. Antes da contratação ou consumo, o Anunciante deve receber preço, modalidade de cobrança e condições essenciais. Entrega depende de saldo, revisão, período, inventário, frequência e regras técnicas, sem promessa de performance.", "Métricas podem incluir entrega, impressão válida, clique, CTR, slot, período e saldo consumido. Relatórios são agregados e a segmentação regional é executada pelo 77Gira sem revelar localização individual ao Anunciante."]),
  ("4. Identificação e medidas", ["Publicidade deve ser identificável na interface. Conteúdo enganoso, ilegal, malicioso ou que viole direitos pode ser removido, e a conta ou campanha pode sofrer medidas proporcionais."])
 ]),
 ("05_regulamento_patacos_e_milipatacos_v1_4.docx", "Regulamento de Patacos e Milipatacos", "Versão 1.4 | Vigência proposta: 30/08/2026", [
  ("1. Natureza e preço", ["Patacos e Milipatacos são unidades internas de mídia para organizar orçamento e veiculação; não são moeda, valor mobiliário, depósito, saque ou conversão automática em dinheiro. Preço, modalidade e condições essenciais são apresentados antes do consumo aplicável."]),
  ("2. Reserva e consumo", ["A campanha usa a fotografia de preço e condições aplicáveis no momento da contratação. Quando houver saldo aplicável, a reserva e o débito seguem o fluxo de campanha e a impressão válida reconhecida pelo sistema. Saldo promocional, se concedido, segue a regra informada na concessão."]),
  ("3. Pagamento real", ["O gateway atual é simulado e não realiza cobrança real nem coleta dados de cartão. Pagamento real, documento fiscal, tributos, reembolso, estorno, chargeback, expiração, transferência e saldo remanescente: " + PENDING_F])
 ]),
 ("06_termo_de_reivindicacao_e_gestao_de_perfil_v1_4.docx", "Termo de Reivindicação e Gestão de Perfil", "Versão 1.4 | Vigência proposta: 30/08/2026", [
  ("1. Legitimidade e aceite", ["Quem solicita ou aceita gestão declara ser titular, integrante autorizado, representante legal ou pessoa com poderes suficientes para agir em nome do Artista, Casa ou organização. O aceite é eletrônico e associado à conta e aos registros disponíveis."]),
  ("2. Análise e gestão", ["O 77Gira pode solicitar informações ou evidências que o fluxo admitir, analisar a solicitação e manter registros de decisão e auditoria. Gestores respondem pela exatidão de agenda, bio, imagens, contatos, links e demais materiais publicados."]),
  ("3. Medidas e contestação", ["Declaração falsa, documento adulterado ou apropriação indevida pode gerar rejeição, revogação, remoção de vínculo, preservação de evidências e medidas cabíveis. Prazo e canal formal de contestação: " + PENDING_J])
 ]),
 ("07_politica_de_conteudo_moderacao_e_denuncias_v1_4.docx", "Política de Conteúdo, Moderação e Denúncias", "Versão 1.4 | Vigência proposta: 30/08/2026", [
  ("1. Conteúdo vedado", ["Não são permitidos fraude, falsa identidade, violação de direitos, assédio, ameaça, discriminação, violência, phishing, link malicioso, exposição indevida de dados, desinformação enganosa ou atividade ilegal."]),
  ("2. Denúncia e análise", ["O 77Gira pode receber denúncia, preservar evidências, pedir esclarecimentos, aplicar correção, limitar alcance ou suspender/retirar conteúdo ou acesso quando proporcional e necessário."]),
  ("3. Urgência", ["Risco grave, malware, fraude, exploração, violência ou obrigação legal pode justificar medida imediata. A plataforma não substitui organizadores, segurança privada, autoridades ou seguro de eventos físicos."])
 ]),
 ("08_politica_de_parcerias_estrategicas_v1_4.docx", "Política de Parcerias Estratégicas", "Versão 1.4 | Vigência proposta: 30/08/2026", [
  ("1. Finalidade", ["Parcerias podem envolver presença institucional, projetos culturais, ativações, conteúdo ou publicidade. Curadoria editorial permanece independente."]),
  ("2. Dados e transparência", ["Relatórios a Parceiros são agregados e não identificam Usuários, salvo base legal e instrumento específico. A presença institucional deve ser identificável quando aplicável."]),
  ("3. Condições específicas", ["Exclusividade, bloqueio de concorrente, investimento, prazo, canais, contrapartidas e uso de marca dependem de instrumento escrito específico."])
 ]),
 ("09_contrato_base_de_parceria_e_patrocinio_v1_4.docx", "Contrato Base de Parceria e Patrocínio", "Minuta-base v1.4 | Revisão jurídica obrigatória antes de assinatura", [
  ("Partes e objeto", ["PARCEIRO: [INFORMAR: razão social, CNPJ, sede e representante]. 77GIRAMUNDO: [RAZÃO SOCIAL DA 77GIRAMUNDO], [CNPJ DA 77GIRAMUNDO]. Objeto, prazo, território, canais, contrapartidas e aprovações devem constar de anexo comercial específico."]),
  ("Marca, dados e independência", ["Cada parte responde pelos dados que tratar. Relatórios são agregados, salvo instrumento e base legal específicos. Marcas permanecem de seus titulares e a licença de uso limita-se ao objeto pactuado. A curadoria editorial do 77Gira permanece independente."]),
  ("Assinatura e pendências", ["Se a adesão ocorrer dentro do 77Gira, o aceite eletrônico seguirá os registros efetivamente mantidos pela plataforma. Se houver assinatura externa, ferramenta, signatários e evidências devem constar deste contrato. Investimento, tributos, rescisão, responsabilidade, confidencialidade e foro: " + PENDING_J])
 ]),
 ("10_contrato_base_de_publicidade_v1_4.docx", "Contrato Base de Publicidade", "Minuta-base v1.4 | Revisão jurídica obrigatória antes de assinatura", [
  ("Partes e campanha", ["ANUNCIANTE: [INFORMAR: razão social, CNPJ, sede e representante]. 77GIRAMUNDO: [RAZÃO SOCIAL DA 77GIRAMUNDO], [CNPJ DA 77GIRAMUNDO]. Campanha, objetivo, período, posicionamentos, criativos, saldo e condições constam do Pedido de Inserção."]),
  ("Veiculação e responsabilidade", ["A veiculação depende de conta aprovada, revisão, saldo, inventário e regras do posicionamento. O Anunciante responde pelo material, oferta, direitos e destino; o 77Gira não garante alcance, vendas, conversão, audiência, posição ou disponibilidade."]),
  ("Métricas e financeiro", ["Relatórios podem apresentar impressão válida, clique, CTR, slot, período e saldo consumido. Preço, impostos, pagamento real, reembolso, estorno, chargeback, rescisão, confidencialidade, responsabilidade e foro dependem de condições comerciais e revisão jurídica: " + PENDING_F])
 ]),
 ("11_termo_de_tratamento_de_dados_e_fornecedores_v1_4.docx", "Termo de Tratamento de Dados e Fornecedores", "Minuta-base v1.4 | Uso interno e contratual", [
  ("1. Papéis", ["Controlador, operador, eventual corresponsável e encarregado devem ser definidos por fluxo e fornecedor. O fornecedor trata dados apenas conforme instruções documentadas, contrato, termos aplicáveis e lei."]),
  ("2. Segurança e transferências", ["O fornecedor deve adotar controles de acesso, proteção de segredos, revisão de permissões, segurança proporcional e apoio em incidentes. Transferências internacionais, suboperadores, países e salvaguardas devem ser confirmados para cada fornecedor antes de publicação contratual."]),
  ("3. Término e pendências", ["No término, dados devem ser devolvidos ou eliminados conforme instruções, ressalvadas retenções legais, segurança, backup ou defesa de direitos. Prazos de backup, retenção, resposta a incidentes e comprovação de eliminação: " + PENDING_J])
 ])
]

def set_font(run, size=11, bold=None, color=None):
    run.font.name = "Calibri"; run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

def style(doc):
    section = doc.sections[0]
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [("Heading 1",16,(46,116,181),18,10),("Heading 2",13,(46,116,181),14,7)]:
        s=doc.styles[name]; s.font.name="Calibri"; s._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); s.font.size=Pt(size); s.font.color.rgb=RGBColor(*color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    p=section.header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=p.add_run("77Gira | Documento jurídico"); set_font(r,9,color=(89,89,89))
    p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("77Giramundo | " + ADDRESS); set_font(r,8,color=(89,89,89))

def build(filename, title, subtitle, sections):
    doc=Document(); style(doc)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); r=p.add_run(title); set_font(r,22,True,(11,37,69))
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(16); r=p.add_run(subtitle); set_font(r,10,color=(89,89,89))
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(14); r=p.add_run(IDENTITY); set_font(r,9,color=(89,89,89))
    for heading, paragraphs in sections:
        doc.add_paragraph(heading, style="Heading 1")
        for text in paragraphs:
            doc.add_paragraph(text)
    doc.add_paragraph("Este documento utiliza aceite eletrônico quando apresentado dentro da plataforma. Itens sinalizados como pendentes exigem definição e revisão jurídica antes de publicação ou assinatura.")
    doc.save(OUT / filename)

for entry in DOCS: build(*entry)
print(f"Gerados {len(DOCS)} documentos em {OUT}")
